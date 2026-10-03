#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.run import Run

ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / "master" / "resume.md"
OUTPUT = ROOT / "output"
OUTPUT.mkdir(exist_ok=True)
DOCX_PATH = OUTPUT / "Mateo_Castano_Staff_Agentic_AI.docx"
PDF_PATH = OUTPUT / "Mateo_Castano_Staff_Agentic_AI.pdf"

BODY_SIZE = 11.5
SUPPORTING_SIZE = 11
LINE_SPACING = 1.05


def font_run(run, size: float, bold: bool = False, italic: bool = False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def keep_with_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    if pPr.find(qn("w:keepNext")) is None:
        pPr.append(OxmlElement("w:keepNext"))


def keep_lines(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    if pPr.find(qn("w:keepLines")) is None:
        pPr.append(OxmlElement("w:keepLines"))


def format_paragraph(p, before=0, after=0, line=LINE_SPACING):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_inline_markdown(p, text: str, size: float, base_bold: bool = False):
    # Supports **bold** inline spans. URLs remain visible plain text for ATS parsing.
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith('**') and part.endswith('**')
        content = part[2:-2] if is_bold else part
        r = p.add_run(content)
        font_run(r, size, bold=(base_bold or is_bold))


def add_contact_links(paragraph, text):
    """Render Markdown contact links as visible, clickable text in normal flow."""
    position = 0
    for match in re.finditer(r"\[([^\]]+)\]\(([^)]+)\)", text):
        font_run(paragraph.add_run(text[position:match.start()]), SUPPORTING_SIZE)
        label, target = match.groups()
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
        element = OxmlElement("w:r")
        hyperlink.append(element)
        run = Run(element, paragraph)
        run.text = label
        font_run(run, SUPPORTING_SIZE)
        run.font.color.rgb = RGBColor.from_string("24527A")
        run.font.underline = True
        paragraph._p.append(hyperlink)
        position = match.end()
    font_run(paragraph.add_run(text[position:]), SUPPORTING_SIZE)


def find_libreoffice():
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable:
        return executable
    mac_executable = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if mac_executable.is_file():
        return str(mac_executable)
    return None


def build_pdf(require_pdf=False):
    soffice = find_libreoffice()
    if not soffice:
        # A previous build's PDF must not appear to match the new DOCX.
        PDF_PATH.unlink(missing_ok=True)
        message = "LibreOffice not found. Install it to generate PDF (macOS: brew install --cask libreoffice)."
        if require_pdf:
            raise SystemExit(message)
        print(f"PDF skipped: {message}")
        return

    PDF_PATH.unlink(missing_ok=True)
    # Isolate the conversion from any open LibreOffice session and stale outputs.
    with tempfile.TemporaryDirectory(prefix="resume-pdf-") as temporary:
        conversion_dir = Path(temporary)
        profile = (conversion_dir / "profile").as_uri()
        try:
            result = subprocess.run(
                [soffice, f"-env:UserInstallation={profile}", "--headless",
                 "--convert-to", "pdf", "--outdir", str(conversion_dir), str(DOCX_PATH)],
                check=True,
                capture_output=True,
                text=True,
                timeout=120,
            )
        except (OSError, subprocess.SubprocessError) as error:
            details = getattr(error, "stderr", None) or str(error)
            raise SystemExit(f"PDF conversion failed: {details}") from error
        generated = conversion_dir / f"{DOCX_PATH.stem}.pdf"
        if not generated.is_file() or generated.stat().st_size == 0:
            raise SystemExit(f"LibreOffice produced no PDF.\n{result.stdout}\n{result.stderr}")
        if not generated.read_bytes().startswith(b"%PDF-"):
            raise SystemExit("LibreOffice output is not a PDF file.")
        shutil.move(str(generated), PDF_PATH)
    print(f"Wrote {PDF_PATH}")


def build_docx(require_pdf=False):
    lines = MASTER.read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.6)
    sec.bottom_margin = Inches(0.6)
    sec.left_margin = Inches(0.65)
    sec.right_margin = Inches(0.65)
    sec.header_distance = Inches(0.2)
    sec.footer_distance = Inches(0.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(BODY_SIZE)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = LINE_SPACING

    first_nonempty_after_h1 = True
    pending_para: list[str] = []
    pending_page_break = False

    def add_paragraph():
        nonlocal pending_page_break
        paragraph = doc.add_paragraph()
        if pending_page_break:
            paragraph.paragraph_format.page_break_before = True
            pending_page_break = False
        return paragraph

    def flush_paragraph():
        nonlocal pending_para
        if not pending_para:
            return
        text = " ".join(x.strip() for x in pending_para).strip()
        p = add_paragraph()
        format_paragraph(p, after=4)
        keep_lines(p)
        add_inline_markdown(p, text, BODY_SIZE)
        pending_para = []

    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            continue

        if stripped == "<!-- PAGEBREAK -->":
            flush_paragraph()
            pending_page_break = True
            continue

        if stripped.startswith("# "):
            flush_paragraph()
            p = add_paragraph()
            format_paragraph(p, after=4)
            r = p.add_run(stripped[2:].upper())
            font_run(r, 22, bold=True)
            first_nonempty_after_h1 = True
            continue

        if stripped.startswith("## "):
            flush_paragraph()
            p = add_paragraph()
            format_paragraph(p, before=9, after=4)
            keep_with_next(p)
            r = p.add_run(stripped[3:].upper())
            font_run(r, 12.5, bold=True)
            continue

        if stripped.startswith("### "):
            flush_paragraph()
            p = add_paragraph()
            format_paragraph(p, before=7, after=3)
            keep_with_next(p)
            add_inline_markdown(p, stripped[4:], BODY_SIZE, base_bold=True)
            continue

        if stripped.startswith("#### "):
            flush_paragraph()
            p = add_paragraph()
            format_paragraph(p, before=6, after=3)
            keep_with_next(p)
            add_inline_markdown(p, stripped[5:], BODY_SIZE, base_bold=True)
            continue

        if stripped.startswith("- "):
            flush_paragraph()
            p = add_paragraph()
            format_paragraph(p, after=3)
            keep_lines(p)
            pf = p.paragraph_format
            pf.left_indent = Inches(0.18)
            pf.first_line_indent = Inches(-0.14)
            r = p.add_run("• ")
            font_run(r, BODY_SIZE)
            add_inline_markdown(p, stripped[2:], BODY_SIZE)
            continue

        # Headline directly after H1.
        if first_nonempty_after_h1 and stripped.startswith("**") and stripped.endswith("**"):
            flush_paragraph()
            p = add_paragraph()
            format_paragraph(p, after=4)
            keep_with_next(p)
            r = p.add_run(stripped[2:-2])
            font_run(r, 12, bold=True)
            first_nonempty_after_h1 = False
            continue

        # Separate contact details and profile URLs into two readable lines.
        if not pending_para and ("@" in stripped or ("linkedin.com/" in stripped and "github.com/" in stripped)):
            flush_paragraph()
            p = add_paragraph()
            is_contact_details = "@" in stripped
            format_paragraph(p, after=2 if is_contact_details else 5)
            keep_lines(p)
            if is_contact_details:
                keep_with_next(p)
            add_contact_links(p, stripped)
            continue

        # Compact bold-label lines used for expertise, skills, education, and earlier experience.
        if stripped.startswith("**") and "**" in stripped[2:]:
            flush_paragraph()
            p = add_paragraph()
            format_paragraph(p, after=3)
            keep_lines(p)
            if stripped.endswith("**") and stripped.count("**") == 2:
                keep_with_next(p)
            add_inline_markdown(p, stripped, SUPPORTING_SIZE)
            continue

        pending_para.append(stripped)
        first_nonempty_after_h1 = False

    flush_paragraph()

    props = doc.core_properties
    props.title = "Mateo Castano - Staff Software Engineer - Agentic AI"
    props.subject = "Resume"
    props.author = "Mateo Castano"
    props.keywords = "Staff Software Engineer, Agentic AI, AI-Native SDLC, Software Modernization, Claude Code, .NET 10, Vue 3"
    doc.save(DOCX_PATH)
    print(f"Wrote {DOCX_PATH}")

    build_pdf(require_pdf=require_pdf)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build the canonical resume as DOCX and optionally PDF.")
    parser.add_argument("--require-pdf", action="store_true", help="Fail if PDF generation is unavailable.")
    args = parser.parse_args()
    build_docx(require_pdf=args.require_pdf)
